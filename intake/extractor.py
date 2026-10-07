"""Turn pasted listing text, an uploaded PDF, or a rent-roll CSV/XLSX into a
draft deal. Claude is only ever asked to pull out fields and flag its own
confidence — it never computes anything. CSV/XLSX rent rolls are parsed with
plain Python, no LLM involved at all.
"""
from __future__ import annotations

import csv
import io
import json
import re

from sqlalchemy.orm import Session

from intake.schemas import IntakeResult
from llm.client import call_claude

_SYSTEM_PROMPT = """You extract real estate deal data from listing text into JSON. \
You never calculate anything (no NOI, cap rate, cash flow, etc.) — only pull out facts that are stated or very strongly implied.

Target JSON shape (fields you can't find: omit them, don't invent values):
{
  "deal": {
    "deal_name": str, "address": str, "price": number, "county": "Lehigh"|"Northampton"|other,
    "property_type": "duplex"|"triplex"|"fourplex"|"mixed_use_residential_majority"|other,
    "units": [{"label": str, "bedrooms": number, "sqft": number, "market_rent": number,
               "current_rent": number|null, "tenant_name": str|null,
               "lease_start": "YYYY-MM-DD"|null, "lease_end": "YYYY-MM-DD"|null,
               "lease_type": "gross"|"NNN"|"modified_gross"|null}],
    "expense_items": [{"name": str, "annual_amount": number, "source": "seller_provided"}],
    "year_built": number|null, "comparable_rent_estimate": number|null
  },
  "confidence": {"<dotted.path.like.price or units[0].market_rent>": "high"|"medium"|"low"},
  "unresolved_notes": ["anything ambiguous, contradictory, or missing that the user should check"]
}

Respond with ONLY that JSON object, no prose, no markdown code fences."""


def _parse_json_response(text: str) -> dict:
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?", "", cleaned).strip()
    cleaned = re.sub(r"```$", "", cleaned).strip()
    return json.loads(cleaned)


def extract_from_text(raw_text: str, session: Session | None = None, deal_id: int | None = None) -> IntakeResult:
    """Extract deal fields from pasted listing text (or text pulled from a PDF)."""
    response_text = call_claude(
        purpose="intake_extraction",
        system=_SYSTEM_PROMPT,
        user_message=raw_text,
        session=session,
        deal_id=deal_id,
        max_tokens=2000,
    )
    try:
        parsed = _parse_json_response(response_text)
    except json.JSONDecodeError as e:
        raise ValueError(f"Claude's extraction response wasn't valid JSON: {e}\n\nRaw response:\n{response_text}") from e

    return IntakeResult(
        deal_dict=parsed.get("deal", {}),
        confidence=parsed.get("confidence", {}),
        unresolved_notes=parsed.get("unresolved_notes", []),
        source_excerpt=raw_text[:2000],
    )


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Pull plain text out of every page of a PDF. No LLM involved."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def extract_from_pdf(file_bytes: bytes, session: Session | None = None, deal_id: int | None = None) -> IntakeResult:
    text = extract_text_from_pdf(file_bytes)
    if not text.strip():
        raise ValueError("Couldn't extract any text from this PDF — it may be a scanned image without OCR.")
    return extract_from_text(text, session=session, deal_id=deal_id)


# ---------------------------------------------------------------------------
# Rent roll CSV / XLSX parsing — pure Python, no LLM
# ---------------------------------------------------------------------------

_COLUMN_SYNONYMS = {
    "label": {"label", "unit", "unit_label", "unit_name"},
    "bedrooms": {"bedrooms", "beds", "bed", "br"},
    "sqft": {"sqft", "sq_ft", "square_feet", "square_footage"},
    "market_rent": {"market_rent", "market rent", "asking_rent", "asking rent"},
    "current_rent": {"current_rent", "current rent", "actual_rent", "actual rent", "rent"},
    "tenant_name": {"tenant", "tenant_name", "tenant name"},
    "lease_start": {"lease_start", "lease start"},
    "lease_end": {"lease_end", "lease end", "lease_expiration", "lease expiration"},
    "lease_type": {"lease_type", "lease type"},
}


def _normalize_header(header: str) -> str | None:
    key = header.strip().lower()
    for field, synonyms in _COLUMN_SYNONYMS.items():
        if key in synonyms:
            return field
    return None


def _coerce_value(field: str, value):
    if value is None or value == "":
        return None
    if field in ("bedrooms", "sqft", "market_rent", "current_rent"):
        try:
            return float(str(value).replace("$", "").replace(",", ""))
        except ValueError:
            return None
    return str(value).strip()


def _normalize_rows(rows: list[dict]) -> list[dict]:
    normalized = []
    for row in rows:
        unit: dict = {}
        for header, value in row.items():
            if header is None:
                continue
            field = _normalize_header(str(header))
            if field is None:
                continue
            coerced = _coerce_value(field, value)
            if coerced is not None:
                unit[field] = coerced
        if unit:
            normalized.append(unit)
    return normalized


def parse_rent_roll_csv(file_bytes: bytes) -> list[dict]:
    """Parse a rent roll CSV into a list of unit dicts, best-effort column matching."""
    text = file_bytes.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return _normalize_rows(list(reader))


def parse_rent_roll_xlsx(file_bytes: bytes) -> list[dict]:
    """Parse a rent roll XLSX (first sheet, first row = header) into unit dicts."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []
    rows = [dict(zip(header, row)) for row in rows_iter]
    return _normalize_rows(rows)
